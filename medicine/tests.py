from io import BytesIO
from unittest.mock import patch

from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User
from django.contrib.messages import get_messages
from django.contrib.messages.storage.cookie import CookieStorage
from django.db import IntegrityError
from django.db.models.deletion import ProtectedError
from django.core.files.uploadedfile import SimpleUploadedFile
from medicine.models import Medicine, MedicineType, MedicineVariant
from medicine.services.medicine_importer import (
    MissingRequiredHeaders,
    import_medicine_rows,
    import_names_only_rows,
    parse_medicine_excel,
    parse_names_only_excel,
    _generate_short_name,
)
from openpyxl import Workbook
import importlib
migration_0008 = importlib.import_module("medicine.migrations.0008_migrate_medicine_types")
OLD_DISPLAY_MAP = migration_0008.OLD_DISPLAY_MAP
SEED_TYPES = migration_0008.SEED_TYPES
migrate_forward = migration_0008.migrate_forward
from django.apps import apps


from doctor.models import InnerMember


class MedicineTypeTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="testdoctor", password="password123")
        InnerMember.objects.create(user=self.user, role="doctor", is_owner=True)
        self.client = Client()
        self.client.login(username="testdoctor", password="password123")

    def test_data_migration_mapping_keeps_types(self):
        """Test that data migration mapping logic correctly maps old choices to MedicineType rows and seeds types."""
        # Verify old choices map to proper title/display names
        expected_mappings = {
            'tablet': 'Tablet',
            'syrup': 'Syrup',
            'injection': 'Injection',
            'capsule': 'Capsule',
            'drops': 'Drops',
            'cream': 'Cream',
        }
        for old_val, expected_display in expected_mappings.items():
            mapped = OLD_DISPLAY_MAP.get(old_val, old_val.title())
            self.assertEqual(mapped, expected_display)

        # Verify seed types list contains all required types
        required_seeds = [
            'Tablet', 'Syrup', 'Injection', 'Capsule', 'Drops', 'Cream',
            'Powder', 'Dry Kadha', 'Liquid Kadha', 'Churna', 'Vati / Gutika',
            'Kwath / Kashayam', 'Asava / Arishta', 'Ghrita', 'Taila (Oil)',
            'Bhasma', 'Avaleha', 'Guggulu', 'Arka', 'Lepa'
        ]
        for seed_name in required_seeds:
            self.assertIn(seed_name, SEED_TYPES)

        # Test executing migrate_forward function
        migrate_forward(apps, None)
        for seed_name in required_seeds:
            self.assertTrue(MedicineType.objects.filter(name=seed_name).exists())

    def _get_messages(self, response):
        cookie = response.cookies.get('messages') or self.client.cookies.get('messages')
        if cookie and cookie.value:
            storage = CookieStorage(response.wsgi_request)
            decoded = storage._decode(cookie.value)
            return [m.message for m in (decoded or [])]
        return [m.message for m in get_messages(response.wsgi_request)]

    def test_used_type_cannot_be_deleted(self):
        """Test that a MedicineType in use by any Medicine cannot be deleted."""
        med_type = MedicineType.objects.create(name="SpecialVati", sort_order=1)
        med = Medicine.objects.create(
            name="Herbal Med",
            short_name="HBM01",
            medicine_type=med_type,
        )

        # Attempt to delete via POST request
        delete_url = reverse("medicine:delete_medicine_type", kwargs={"pk": med_type.pk})
        response = self.client.post(delete_url)

        # Type should NOT be deleted
        self.assertTrue(MedicineType.objects.filter(pk=med_type.pk).exists())

        # Error message should be in messages
        messages = self._get_messages(response)
        self.assertIn("This type is used by 1 medicines. Deactivate it instead.", messages)

        # Model PROTECT should also block direct deletion
        with self.assertRaises(ProtectedError):
            med_type.delete()

        # Unused type can be deleted successfully
        unused_type = MedicineType.objects.create(name="UnusedType", sort_order=2)
        del_unused_url = reverse("medicine:delete_medicine_type", kwargs={"pk": unused_type.pk})
        response2 = self.client.post(del_unused_url)
        self.assertFalse(MedicineType.objects.filter(pk=unused_type.pk).exists())

    def test_duplicate_name_rejected(self):
        """Test that duplicate MedicineType names (case-insensitive) are rejected."""
        MedicineType.objects.create(name="TabletForm")

        # Duplicate via view POST (case-insensitive)
        url = reverse("medicine:medicine_types")
        response = self.client.post(url, {"name": "tabletform", "sort_order": "0", "is_active": "on"})
        messages = self._get_messages(response)
        self.assertIn("Medicine type 'tabletform' already exists.", messages)
        self.assertEqual(MedicineType.objects.filter(name__iexact="tabletform").count(), 1)

        # Duplicate at model level
        with self.assertRaises(IntegrityError):
            MedicineType.objects.create(name="TabletForm")


class MedicineImporterTests(TestCase):
    def _workbook(self, headers, values):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        sheet.append(values)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        return output

    def _upload_file(self, headers, values):
        workbook = self._workbook(headers, values)
        return SimpleUploadedFile(
            'medicines.xlsx',
            workbook.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    def _login_import_user(self):
        user = User.objects.create_user(username='medicine-import-doctor', password='password123')
        InnerMember.objects.create(user=user, role='doctor', is_owner=True)
        self.client.login(username='medicine-import-doctor', password='password123')

    def _row(self, short_name, name=None, packs=3):
        return {
            'row_num': 2,
            'name': name or short_name,
            'short_name': short_name,
            'type': '',
            'company': '',
            'pack_size': '100 gm',
            'unit': 'jar',
            'units_per_pack': 2,
            'cost_price': '10.00',
            'selling_price': '15.00',
            'packs': packs,
            'stock_units': packs * 2,
            'low_stock_alert': 10,
            'mfg_date': '',
            'exp_date': '',
            'status': 'Update existing',
            'reason': '',
        }

    def test_headers_match_aliases_case_insensitively_in_any_order(self):
        workbook = self._workbook(
            [' MANUFACTURER ', 'Code', 'Selling Price', 'Strength', 'Medicine Name', 'COST_PRICE', 'Medicine Type', 'Unit', 'stock_strips', 'unit_per_strip'],
            ['Ayur Care', 'ASH-01', 120, '100 gm', 'Ashwagandha Churna', 80, 'Churna', 'JAR', 5, 1],
        )

        rows = parse_medicine_excel(workbook)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['name'], 'Ashwagandha Churna')
        self.assertEqual(rows[0]['short_name'], 'ASH-01')
        self.assertEqual(rows[0]['pack_size'], '100 gm')
        self.assertEqual(rows[0]['unit'], 'jar')
        self.assertEqual(rows[0]['stock_units'], 5)

    def test_missing_required_header_lists_expected_header_row(self):
        workbook = self._workbook(
            ['name', 'short_name', 'pack_size', 'selling_price'],
            ['Ashwagandha', 'ASH-02', '100 gm', 100],
        )

        with self.assertRaises(MissingRequiredHeaders) as raised:
            parse_medicine_excel(workbook)

        self.assertIn('cost_price', str(raised.exception))
        self.assertIn('Expected header row:', str(raised.exception))

    def test_zero_selling_price_is_rejected_in_preview(self):
        workbook = self._workbook(
            ['name', 'short_name', 'pack_size', 'cost_price', 'selling_price'],
            ['Ashwagandha', 'ASH-03', '100 gm', 50, 0],
        )

        row = parse_medicine_excel(workbook)[0]

        self.assertEqual(row['status'], 'Error')
        self.assertIn('selling price 0', row['reason'])

    def test_blank_units_per_pack_treats_packs_as_stock_units(self):
        workbook = self._workbook(
            ['name', 'short_name', 'pack_size', 'cost_price', 'selling_price', 'packs', 'units_per_pack'],
            ['Ashwagandha', 'ASH-04', '100 gm', 50, 75, 4, None],
        )

        row = parse_medicine_excel(workbook)[0]

        self.assertEqual(row['status'], 'New medicine')
        self.assertIsNone(row['units_per_pack'])
        self.assertEqual(row['stock_units'], 4)

    def test_add_and_replace_stock_options(self):
        medicine = Medicine.objects.create(name='Test medicine', short_name='TEST-01')
        variant = MedicineVariant.objects.create(
            medicine=medicine,
            power='100 gm',
            cost_price=10,
            selling_price=15,
            stock=20,
        )
        row = self._row('TEST-01')

        import_medicine_rows([row], 'add')
        variant.refresh_from_db()
        self.assertEqual(variant.stock, 26)
        import_medicine_rows([row], 'replace')
        variant.refresh_from_db()
        self.assertEqual(variant.stock, 6)
        import_medicine_rows([row], 'keep')
        variant.refresh_from_db()
        self.assertEqual(variant.stock, 6)

    def test_one_database_error_does_not_rollback_other_rows(self):
        first_row = self._row('ISOLATE-01')
        second_row = self._row('ISOLATE-02')
        actual_create = MedicineVariant.objects.create
        call_count = 0

        def fail_first_create(**kwargs):
            nonlocal call_count
            call_count += 1
            if call_count == 1:
                raise RuntimeError('simulated row failure')
            return actual_create(**kwargs)

        with patch('medicine.services.medicine_importer.MedicineVariant.objects.create', side_effect=fail_first_create):
            result = import_medicine_rows([first_row, second_row], 'add')

        self.assertEqual(result['success_count'], 1)
        self.assertEqual(len(result['error_rows']), 1)
        self.assertFalse(Medicine.objects.filter(short_name='ISOLATE-01').exists())
        self.assertTrue(Medicine.objects.filter(short_name='ISOLATE-02').exists())

    def test_import_preview_requires_confirmation_when_adding_same_file_twice(self):
        self._login_import_user()
        headers = ['name', 'short_name', 'pack_size', 'cost_price', 'selling_price', 'packs', 'units_per_pack']
        values = ['Ashwagandha Churna', 'ASH-RE', '100 gm', 50, 75, 3, 2]
        url = reverse('medicine:import_medicine')
        file_content = self._upload_file(headers, values).read()

        def same_file_upload():
            return SimpleUploadedFile(
                'medicines.xlsx',
                file_content,
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )

        first_preview = self.client.post(url, {
            'action': 'upload',
            'stock_option': 'add',
            'excel_file': same_file_upload(),
        })
        self.assertEqual(first_preview.status_code, 200)
        self.assertFalse(Medicine.objects.filter(short_name='ASH-RE').exists())
        self.client.post(url, {'action': 'import'})
        variant = MedicineVariant.objects.get(medicine__short_name='ASH-RE')
        self.assertEqual(variant.stock, 6)

        second_preview = self.client.post(url, {
            'action': 'upload',
            'stock_option': 'add',
            'excel_file': same_file_upload(),
        })
        self.assertContains(second_preview, 'Importing again will add the stock twice.')
        blocked_import = self.client.post(url, {'action': 'import'})
        self.assertContains(blocked_import, 'I confirm adding this file again.')
        variant.refresh_from_db()
        self.assertEqual(variant.stock, 6)

        self.client.post(url, {'action': 'import', 'confirm_duplicate': 'on'})
        variant.refresh_from_db()
        self.assertEqual(variant.stock, 12)

    def test_names_only_client_layout_parses_sections_and_headerless_company_on_all_sheets(self):
        workbook = Workbook()
        first = workbook.active
        first.title = 'Sheet1'
        first.append(['Client inventory'])
        first.append(['S.No', 'MEDICINE NAME', None])
        first.append([None, 'CHURNA', None])
        first.append([1, 'Mahasudarshan Churana', 'KALEDA'])
        second = workbook.create_sheet('Taila')
        second.append([None, 'Medicine Name', 'Company'])
        second.append([1, 'Mahanarayan Taila', 'LOCAL'])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)

        rows = parse_names_only_excel(output)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['sheet_name'], 'Sheet1')
        self.assertEqual(rows[0]['type'], 'Churna')
        self.assertEqual(rows[0]['name'], 'Mahasudarshan Churna')
        self.assertEqual(rows[0]['company'], 'KALEDA')
        self.assertEqual(rows[0]['short_name'], 'MC')
        self.assertEqual(rows[1]['type'], 'Taila (Oil)')
        self.assertEqual(rows[1]['company'], 'LOCAL')

    def test_names_only_short_names_are_unique_and_single_word_uses_four_letters(self):
        used = set()

        self.assertEqual(_generate_short_name('Mahasudarshan Churna', used), 'MC')
        self.assertEqual(_generate_short_name('Mahanarayan Churna', used), 'MC2')
        self.assertEqual(_generate_short_name('Sunth', used), 'SUNT')

    def test_unknown_type_is_previewed_and_created_in_title_case(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['S.No', 'MEDICINE NAME', 'TYPE', 'COMPANY'])
        sheet.append([1, 'Neeri Tablet', '  NEERI TABLET  ', 'LOCAL'])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)

        rows = parse_names_only_excel(output)

        self.assertEqual(rows[0]['type'], 'Neeri Tablet')
        self.assertEqual(rows[0]['type_status'], 'New type will be created: Neeri Tablet')
        result = import_names_only_rows(rows)

        self.assertEqual(result['success_count'], 1)
        self.assertTrue(MedicineType.objects.filter(name='Neeri Tablet').exists())
        medicine = Medicine.objects.get(name='Neeri Tablet')
        self.assertEqual(medicine.medicine_type.name, 'Neeri Tablet')

    def test_names_only_duplicates_are_skipped_and_import_creates_no_variants(self):
        Medicine.objects.create(name='Ashwagandha Churna', short_name='OLD-AS')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Sheet1'
        sheet.append(['S.No', 'MEDICINE NAME', None])
        sheet.append([1, 'Ashwagandha Churna', 'OLD'])
        sheet.append([2, 'Neem Churna', 'LOCAL'])
        sheet.append([3, 'Neem Churna', 'LOCAL'])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)

        rows = parse_names_only_excel(output)

        self.assertEqual([row['status'] for row in rows], ['Already exists', 'New', 'Already exists'])
        result = import_names_only_rows(rows)
        self.assertEqual(result['success_count'], 1)
        self.assertEqual(Medicine.objects.filter(name__iexact='Neem Churna').count(), 1)
        neem = Medicine.objects.get(name__iexact='Neem Churna')
        self.assertFalse(neem.variants.exists())

    def test_variantless_medicine_is_flagged_in_store_and_grouped_search(self):
        self._login_import_user()
        incomplete = Medicine.objects.create(name='Sunth', short_name='SUNT')

        store_response = self.client.get(reverse('medicine:manage_medicine'))
        self.assertContains(store_response, 'Add pack &amp; price')
        self.assertContains(store_response, '<option value="incomplete">Incomplete</option>', html=True)
        self.assertNotContains(store_response, 'Incomplete (1)')

        filtered_response = self.client.get(reverse('medicine:manage_medicine'), {'pack': 'incomplete'})
        self.assertEqual(
            [medicine.pk for medicine in filtered_response.context['medicines'].object_list],
            [incomplete.pk],
        )

        search_response = self.client.get(
            reverse('medicine:search_medicine'),
            {'q': 'Sunth', 'grouped': '1'},
        )
        medicine_result = search_response.json()['medicines'][0]
        self.assertEqual(medicine_result['id'], incomplete.pk)
        self.assertTrue(medicine_result['price_not_set'])
        self.assertEqual(medicine_result['variants'], [])

    def test_manage_medicine_filters_and_query_context(self):
        self._login_import_user()
        med_type = MedicineType.objects.create(name='Filter type')
        in_stock = Medicine.objects.create(
            name='In stock medicine', short_name='IN-STOCK', medicine_type=med_type,
            company='Filter Co', is_active=True,
        )
        low_stock = Medicine.objects.create(name='Low stock medicine', short_name='LOW-STOCK')
        out_of_stock = Medicine.objects.create(name='Out of stock medicine', short_name='OUT-STOCK')
        for medicine, stock in ((in_stock, 20), (low_stock, 5), (out_of_stock, 0)):
            MedicineVariant.objects.create(
                medicine=medicine, power='1 pack', cost_price=10, selling_price=15,
                stock=stock, low_stock_alert=10,
            )

        url = reverse('medicine:manage_medicine')
        response = self.client.get(url, {
            'type': str(med_type.pk), 'company': 'Filter Co', 'pack': 'complete',
            'stock': 'in_stock', 'status': 'active',
        })
        self.assertEqual(
            [medicine.pk for medicine in response.context['medicines'].object_list],
            [in_stock.pk],
        )
        self.assertEqual(response.context['active_filter_count'], 5)
        self.assertIn('type=', response.context['pagination_query'])
        self.assertIn('company=Filter+Co', response.context['pagination_query'])

        for stock_filter, expected in (
            ('low', low_stock), ('out', out_of_stock),
        ):
            stock_response = self.client.get(url, {'stock': stock_filter})
            self.assertEqual(
                [medicine.pk for medicine in stock_response.context['medicines'].object_list],
                [expected.pk],
            )

